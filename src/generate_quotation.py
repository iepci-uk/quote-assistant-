"""Fill the IEPCI quotation template from an RFQ (draft, prices not approved).

Run from the project root:
    python -m src.generate_quotation

The Streamlit app (app.py) calls generate_quotation() with the lines and RFQ
details a person has reviewed.
"""

from copy import copy
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils.cell import range_boundaries


PROJECT_ROOT = Path(__file__).resolve().parent.parent

TEMPLATE_PATH = PROJECT_ROOT / "data" / "raw" / "IEPCI Quotation .xlsx"
RFQ_PATH = PROJECT_ROOT / "data" / "raw" / "RFQ_KBS-0147.pdf"
OUTPUT_DIR = PROJECT_ROOT / "data" / "processed"

FIRST_ITEM_ROW = 13
LAST_TEMPLATE_ITEM_ROW = 29
CHARS_PER_DESCRIPTION_LINE = 46   # merged C:H cell at the template's 16 pt font


def copy_row_format(ws, source_row, target_row):
    """Copy cell formatting and row height to a new item row."""
    for col in range(1, ws.max_column + 1):
        source = ws.cell(source_row, col)
        target = ws.cell(target_row, col)

        if source.has_style:
            target._style = copy(source._style)

        if source.number_format:
            target.number_format = source.number_format

        if source.alignment:
            target.alignment = copy(source.alignment)

        if source.protection:
            target.protection = copy(source.protection)

    height = ws.row_dimensions[source_row].height
    if height is not None:
        ws.row_dimensions[target_row].height = height


def expand_item_table(ws, item_count):
    """Insert item rows and correctly restore merged cells below the table."""
    capacity = LAST_TEMPLATE_ITEM_ROW - FIRST_ITEM_ROW + 1

    if item_count <= capacity:
        return None

    extra_rows = item_count - capacity
    insert_at = LAST_TEMPLATE_ITEM_ROW + 1

    # Save and temporarily remove merges that could be affected by insertion.
    affected_merges = [
        str(merged_range)
        for merged_range in list(ws.merged_cells.ranges)
        if merged_range.max_row >= insert_at
    ]

    for merged_range in affected_merges:
        ws.unmerge_cells(merged_range)

    # openpyxl does not automatically move row heights when inserting rows.
    saved_heights = {
        row: ws.row_dimensions[row].height
        for row in list(ws.row_dimensions)
        if row >= insert_at
    }

    for row in saved_heights:
        del ws.row_dimensions[row]

    ws.insert_rows(insert_at, extra_rows)

    # openpyxl does not move pictures when rows are inserted. Move the ones
    # below the item table (the company stamp) down with their cells.
    for image in ws._images:
        anchor = image.anchor
        if anchor._from.row >= insert_at - 1:          # 0-based row index
            anchor._from.row += extra_rows
            if getattr(anchor, "to", None) is not None:
                anchor.to.row += extra_rows

    # Restore row heights at their new positions.
    for row, height in saved_heights.items():
        if height is not None:
            ws.row_dimensions[row + extra_rows].height = height

    # Format the additional item rows using the final template item row.
    for row in range(insert_at, insert_at + extra_rows):
        copy_row_format(ws, LAST_TEMPLATE_ITEM_ROW, row)

    # Restore the original merged cells at their shifted positions.
    for merged_range in affected_merges:
        min_col, min_row, max_col, max_row = range_boundaries(
            merged_range
        )

        if min_row >= insert_at:
            min_row += extra_rows
            max_row += extra_rows
        elif max_row >= insert_at:
            max_row += extra_rows

        ws.merge_cells(
            start_row=min_row,
            start_column=min_col,
            end_row=max_row,
            end_column=max_col,
        )

    # Create the usual item-row merges for each additional row.
    for row in range(insert_at, insert_at + extra_rows):
        for start_col, end_col in [
            (3, 8),    # C:H — Description
            (11, 12),  # K:L — Unit rate
            (13, 14),  # M:N — Line total
        ]:
            ws.merge_cells(
                start_row=row,
                start_column=start_col,
                end_row=row,
                end_column=end_col,
            )

    last_item_row = FIRST_ITEM_ROW + item_count - 1
    subtotal_row = last_item_row + 1
    tax_row = subtotal_row + 1
    packing_row = subtotal_row + 2
    freight_row = subtotal_row + 3
    inspection_row = subtotal_row + 4
    total_row = subtotal_row + 5

    return {
        "last_item_row": last_item_row,
        "subtotal_row": subtotal_row,
        "tax_row": tax_row,
        "packing_row": packing_row,
        "freight_row": freight_row,
        "inspection_row": inspection_row,
        "total_row": total_row,
    }


def fit_item_row(ws, row):
    """Wrap a long description and make the row tall enough to show it.

    The template uses a 16 pt font in a merged C:H cell (about 46 characters
    per line) and a fixed 25 pt row height, so long RFQ descriptions were cut
    off. Rows stay at the template height when the text fits on one line.
    """
    text = str(ws.cell(row, 3).value or "")
    lines = max(1, -(-len(text) // CHARS_PER_DESCRIPTION_LINE))
    ws.row_dimensions[row].height = max(25, 21 * lines + 4)

    description = ws.cell(row, 3)
    description.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    for column in (1, 2, 9, 10, 11, 13):
        cell = ws.cell(row, column)
        alignment = copy(cell.alignment)
        alignment.vertical = "center"
        if column == 1:
            alignment.horizontal = "center"         # keep Sr No apart from the code
        if column == 2:
            alignment.shrink_to_fit = True          # long part codes
        cell.alignment = alignment

    ws.cell(row, 11).number_format = "#,##0.00"     # unit rate, like the totals


def split_address(address):
    """Split one address string over the template's two address lines."""
    parts = [p.strip() for p in str(address or "").split(",") if p.strip()]
    if len(parts) <= 1:
        return (parts[0] if parts else ""), ""
    middle = (len(parts) + 1) // 2
    return ", ".join(parts[:middle]), ", ".join(parts[middle:])


def fill_header(ws, header):
    """Write the customer / RFQ details into the template's header block.

    Template cells (merged ranges, so we write to the top-left cell):
      C6 Company name   C7 Address line 1   C8 Address line 2
      C9 Phone          C10 Email
      M7 Quotation date (header['quote_date'], default today)
      M8 Inquiry ref      M9 Inquiry date
    Empty values are left blank. M6 (our own quote number), M10 (inquiry
    type), M11 (sales representative), GSTN and IEC are not in the RFQ.
    """
    header = header or {}
    line1, line2 = split_address(header.get("address"))

    values = {
        "C6": header.get("customer_name"),
        "C7": line1,
        "C8": line2,
        "C9": header.get("phone"),
        "C10": header.get("email"),
        "M7": header.get("quote_date") or date.today().strftime("%d/%m/%Y"),
        "M8": header.get("rfq_number"),
        "M9": header.get("rfq_date"),
    }
    for cell, value in values.items():
        if value:
            ws[cell] = str(value)


def generate_quotation(rfq_path=RFQ_PATH, output_path=None, rfq_df=None, header=None, use_list_prices=True):
    """Create a draft Excel quotation from an RFQ and the IEPCI template.

    header: dict of RFQ details (see extract_rfq.HEADER_FIELDS), usually the
    values a person checked in the app. If it is not given and the lines are
    read from rfq_path, the header is read from that file too.

    use_list_prices: when the lines are read from rfq_path (rfq_df not given),
    fill the rate with the supplier list price (INR, no discount) wherever the
    price is clear. Lines without a clear price stay blank and say REVIEW.
    """
    from src.extract_rfq import extract_rfq_header, extract_rfq_lines

    if rfq_df is None:
        rfq_df = extract_rfq_lines(str(rfq_path))
        if header is None:
            header = extract_rfq_header(str(rfq_path))
        if use_list_prices:
            from src.load_price_lists import load_catalogue
            from src.pricing import apply_list_prices

            rfq_df = apply_list_prices(rfq_df, load_catalogue())

    if rfq_df.empty:
        raise ValueError("The RFQ contains no extracted line items.")

    if len(rfq_df) > 100:
        raise ValueError(
            "This RFQ has more than 100 items; review it before generating."
        )

    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(
            f"Quotation template not found: {TEMPLATE_PATH}"
        )

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    if output_path is None:
        output_path = OUTPUT_DIR / "IEPCI_Quotation_DRAFT.xlsx"
    else:
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = load_workbook(TEMPLATE_PATH)
    ws = workbook["Format"]

    layout = expand_item_table(ws, len(rfq_df))

    if layout is None:
        last_item_row = FIRST_ITEM_ROW + len(rfq_df) - 1
        subtotal_row = LAST_TEMPLATE_ITEM_ROW + 1
        tax_row = subtotal_row + 1
        packing_row = subtotal_row + 2
        freight_row = subtotal_row + 3
        inspection_row = subtotal_row + 4
        total_row = subtotal_row + 5
    else:
        last_item_row = layout["last_item_row"]
        subtotal_row = layout["subtotal_row"]
        tax_row = layout["tax_row"]
        packing_row = layout["packing_row"]
        freight_row = layout["freight_row"]
        inspection_row = layout["inspection_row"]
        total_row = layout["total_row"]

    # Fill only the RFQ details that are known.
    ws["M6"] = "IEPCI/QTN/2026/DRAFT-01"
    fill_header(ws, header)

    # All amounts are in INR (the template and the supplier price lists).
    ws["A5"] = "QUOTATION - DRAFT (PRICES TO BE VERIFIED)"

    # Populate all extracted RFQ items.
    for index, (_, item) in enumerate(rfq_df.iterrows()):
        row = FIRST_ITEM_ROW + index

        line_no = item.get("line_no", index + 1)
        make = str(item.get("make", "") or "").strip()
        part_code = str(item.get("part_code", "") or "").strip()
        description = str(item.get("description", "") or "").strip()
        code_note = str(item.get("code_note", "") or "").strip()
        qty = item.get("qty", None)
        unit = str(item.get("unit", "") or "").strip()

        ws.cell(row, 1).value = (
            int(line_no) if str(line_no).isdigit() else index + 1
        )
        ws.cell(row, 2).value = part_code or None

        description_parts = []
        if make:
            description_parts.append(make)
        if description:
            description_parts.append(description)
        if code_note:
            description_parts.append(f"REVIEW: {code_note}")

        ws.cell(row, 3).value = " — ".join(description_parts)
        fit_item_row(ws, row)
        ws.cell(row, 9).value = unit
        ws.cell(row, 10).value = int(qty) if str(qty).isdigit() else qty

        # No unapproved selling price is entered.
        unit_rate = item.get("unit_rate")
        approved = item.get("price_approved", False)

        if approved and unit_rate is not None and str(unit_rate).strip().lower() not in ("", "nan"):
            ws.cell(row, 11).value = float(unit_rate)
        else:
            ws.cell(row, 11).value = None

        # Line total for EVERY item row. (This used to sit inside the
        # else-block above, so approved rows beyond the template's 17 rows
        # never got a total and were left out of the subtotal.)
        ws.cell(row, 13).value = (
            f'=IF(OR(K{row}="",J{row}=""),"",K{row}*J{row})'
        )

    # Keep labels aligned with the template's commercial section.
    ws.cell(subtotal_row, 3).value = "Subtotal"
    ws.cell(tax_row, 3).value = "Taxes"
    ws.cell(packing_row, 3).value = "Packing and Forwarding"
    ws.cell(freight_row, 3).value = "Freight and Insurance"
    ws.cell(inspection_row, 3).value = "Inspection"
    ws.cell(total_row, 3).value = "Total"

    # Calculate totals only when rates have been entered.
    ws.cell(subtotal_row, 13).value = (
        f'=IF(COUNT(K{FIRST_ITEM_ROW}:K{last_item_row})=0,"",'
        f'SUM(M{FIRST_ITEM_ROW}:M{last_item_row}))'
    )

    # The template uses 18% GST; applicability must be confirmed before issue.
    ws.cell(tax_row, 10).value = 0.18
    ws.cell(tax_row, 13).value = (
        f'=IF(M{subtotal_row}="","",M{subtotal_row}*18%)'
    )

    # Packing is a direct amount; leave it blank until approved.
    ws.cell(packing_row, 13).value = None

    ws.cell(freight_row, 13).value = (
        f'=IF(OR(K{freight_row}="",J{freight_row}=""),"",'
        f'K{freight_row}*J{freight_row})'
    )

    ws.cell(inspection_row, 13).value = (
        f'=IF(OR(K{inspection_row}="",J{inspection_row}=""),"",'
        f'K{inspection_row}*J{inspection_row})'
    )

    ws.cell(total_row, 13).value = (
        f'=IF(M{subtotal_row}="","",'
        f'M{subtotal_row}+M{tax_row}+N(M{packing_row})'
        f'+N(M{freight_row})+N(M{inspection_row}))'
    )

    # Keep the draft warning in the title instead of overwriting row 36.
    # Row 36 and the following rows remain part of the quotation layout.

    try:
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.calculation.calcMode = "auto"
    except AttributeError:
        pass

    workbook.save(output_path)
    return output_path


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="Create a draft Excel quotation from an RFQ.")
    parser.add_argument("rfq", nargs="?", default=RFQ_PATH, help="RFQ file (PDF, DOCX or XLSX)")
    parser.add_argument("--out", help="where to save the quotation (default: data/processed/)")
    parser.add_argument("--date", help="quotation date as dd/mm/yyyy (default: today)")
    parser.add_argument(
        "--list-prices",
        action="store_true",
        help="use the supplier list prices (INR, no discount) as the rates",
    )
    args = parser.parse_args()

    from src.extract_rfq import extract_rfq_header, extract_rfq_lines

    lines = extract_rfq_lines(args.rfq)
    if args.list_prices:
        from src.load_price_lists import load_catalogue
        from src.pricing import apply_list_prices

        lines = apply_list_prices(lines, load_catalogue())

    header = extract_rfq_header(args.rfq)
    if args.date:
        header["quote_date"] = args.date
    path = generate_quotation(output_path=args.out, rfq_df=lines, header=header)
    print(f"Draft quotation created: {path}")
    if args.list_prices:
        priced = int(lines["price_approved"].astype(bool).sum())
        print(f"{priced} of {len(lines)} lines priced at supplier list price (INR, no discount).")
        print("Lines without a clear price are blank and marked REVIEW.")
    else:
        print("Unit rates and final totals have not been approved.")
